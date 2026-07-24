from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


GREEN = RGBColor(23, 75, 58)
ACCENT = RGBColor(45, 125, 97)
LIGHT = RGBColor(238, 245, 242)
DARK = RGBColor(30, 43, 38)
MUTED = RGBColor(91, 108, 101)
WHITE = RGBColor(255, 255, 255)
WARN = RGBColor(194, 116, 37)
RED = RGBColor(175, 56, 56)


class MeetingPptxRenderer:
    """Build an editable 16:9 commercial deck from a persisted report snapshot."""

    def render(self, report) -> bytes:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        payload = report["payload"]
        self._cover(presentation, report)
        self._executive(presentation, report)
        self._ranking_chart(presentation, payload.get("ranking", []), report)
        self._ranking_table(presentation, payload.get("ranking", []), report)
        self._objectives(presentation, payload.get("objectives", []), report)
        self._alerts(presentation, payload.get("alerts", []), report)
        self._recognitions(presentation, payload.get("recognitions", []), report)
        if payload.get("seller_pages"):
            self._seller_overview(presentation, payload["seller_pages"], report)
        self._commitments(presentation, report)
        if len(presentation.slides) > 12:
            raise ValueError("La presentación excede el máximo de 12 diapositivas")
        buffer = BytesIO()
        presentation.save(buffer)
        return buffer.getvalue()

    def _cover(self, prs, report):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._background(slide, GREEN)
        self._textbox(slide, 0.8, 1.4, 11.7, 0.7, "CODENOA SALES COACH", 16, WHITE, bold=True)
        self._textbox(slide, 0.8, 2.2, 11.7, 1.2, "Reunión comercial", 34, WHITE, bold=True)
        period = report["period"]
        self._textbox(
            slide, 0.8, 3.55, 11.7, 0.5,
            f"Período {period['start']} a {period['end']}", 18, WHITE,
        )
        self._textbox(
            slide, 0.8, 5.8, 11.7, 0.5,
            f"Snapshot {report['report_id'][:8]} · versión {report['version']} · "
            f"actualización {report.get('data_updated_at') or 'sin dato'}",
            9, WHITE,
        )

    def _executive(self, prs, report):
        payload = report["payload"]
        slide = self._slide(prs, "Resumen ejecutivo", report)
        kpis = payload["kpis"]
        values = [
            ("Venta neta", self._money(kpis.get("net_sales"))),
            ("Cantidad", self._number(kpis.get("quantity"))),
            ("Clientes activos", self._number(kpis.get("active_clients"))),
            ("Crecimiento", f"{self._number(kpis.get('growth_pct'))}%"),
            ("Cumplimiento", f"{self._number(kpis.get('objective_fulfillment_pct'))}%"),
            ("Potencial", self._money(kpis.get("opportunity_potential"))),
        ]
        for index, (label, value) in enumerate(values):
            col, row = index % 3, index // 3
            self._card(slide, 0.65 + col * 4.15, 1.25 + row * 1.25, 3.8, 0.95, label, value)
        self._bullet_box(slide, 0.65, 4.05, 12.0, 2.25, payload.get("executive_summary", []), "Lectura del período")

    def _ranking_chart(self, prs, ranking, report):
        slide = self._slide(prs, "Desempeño comercial", report)
        rows = ranking[:10]
        if not rows:
            self._empty(slide, "No hay vendedores con datos para graficar.")
            return
        chart_data = ChartData()
        chart_data.categories = [str(row.get("seller") or "")[:18] for row in rows]
        chart_data.add_series("Venta neta", [float(row.get("sales") or 0) for row in rows])
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(0.7), Inches(1.25), Inches(8.2), Inches(5.45),
            chart_data,
        ).chart
        chart.has_legend = False
        chart.value_axis.has_major_gridlines = True
        chart.category_axis.reverse_order = True
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = ACCENT
        growth = sorted(ranking, key=lambda row: float(row.get("growthPct") or 0), reverse=True)[:5]
        self._bullet_box(
            slide, 9.2, 1.25, 3.45, 5.45,
            [
                f"{row.get('seller')}: {float(row.get('growthPct') or 0):.1f}%"
                for row in growth
            ],
            "Mayor crecimiento",
        )

    def _ranking_table(self, prs, ranking, report):
        slide = self._slide(prs, "Ranking de vendedores", report)
        rows = ranking[:10]
        table = slide.shapes.add_table(
            len(rows) + 1, 6, Inches(0.55), Inches(1.25), Inches(12.25), Inches(5.6)
        ).table
        widths = [0.6, 3.2, 2.25, 1.65, 1.7, 1.4]
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        headers = ["#", "Vendedor", "Venta", "Crecimiento", "Clientes", "Mix"]
        for col, label in enumerate(headers):
            self._table_cell(table.cell(0, col), label, GREEN, WHITE, True)
        for row_index, row in enumerate(rows, start=1):
            values = [
                row.get("rankSales", ""),
                row.get("seller", ""),
                self._money(row.get("sales")),
                f"{self._number(row.get('growthPct'))}%",
                self._number(row.get("clients")),
                self._number(row.get("mixCount")),
            ]
            for col, value in enumerate(values):
                self._table_cell(
                    table.cell(row_index, col), str(value),
                    LIGHT if row_index % 2 else WHITE, DARK, col == 1,
                )

    def _objectives(self, prs, objectives, report):
        slide = self._slide(prs, "Objetivos comerciales", report)
        rows = objectives[:8]
        if not rows:
            self._empty(slide, "No hay objetivos activos para el período.")
            return
        table = slide.shapes.add_table(
            len(rows) + 1, 5, Inches(0.65), Inches(1.35), Inches(12.0), Inches(5.25)
        ).table
        headers = ["Alcance", "Métrica", "Real", "Objetivo", "Cumplimiento"]
        for col, label in enumerate(headers):
            self._table_cell(table.cell(0, col), label, GREEN, WHITE, True)
        for row_index, objective in enumerate(rows, start=1):
            progress = objective.get("progress") or {}
            values = [
                objective.get("scope_key", ""),
                objective.get("metric", ""),
                progress.get("actual_value", 0),
                objective.get("target_value", 0),
                f"{float(progress.get('fulfillment_pct') or 0):.1f}%",
            ]
            for col, value in enumerate(values):
                self._table_cell(table.cell(row_index, col), str(value), LIGHT if row_index % 2 else WHITE, DARK)

    def _alerts(self, prs, alerts, report):
        slide = self._slide(prs, "Riesgos y oportunidades", report)
        rows = alerts[:8]
        if not rows:
            self._empty(slide, "No hay alertas abiertas priorizadas para este corte.")
            return
        for index, alert in enumerate(rows):
            col, row = index % 2, index // 2
            severity = alert.get("severity")
            color = RED if severity == "critical" else WARN if severity == "warning" else ACCENT
            self._card(
                slide, 0.65 + col * 6.05, 1.25 + row * 1.35, 5.7, 1.05,
                str(alert.get("type") or "Alerta"),
                str(alert.get("message") or "")[:130],
                accent=color,
                value_size=11,
            )

    def _recognitions(self, prs, recognitions, report):
        slide = self._slide(prs, "Reconocimientos del período", report)
        if not recognitions:
            self._empty(slide, "No se detectaron reconocimientos bajo las reglas vigentes.")
            return
        for index, item in enumerate(recognitions[:6]):
            self._card(
                slide, 0.75, 1.25 + index * 0.85, 11.8, 0.7,
                str(item.get("seller") or ""), str(item.get("message") or ""),
                accent=ACCENT, value_size=12,
            )

    def _seller_overview(self, prs, pages, report):
        slide = self._slide(prs, "Focos de coaching por vendedor", report)
        selected = sorted(
            pages,
            key=lambda item: float(item.get("kpis", {}).get("sales") or 0),
            reverse=True,
        )[:6]
        for index, seller in enumerate(selected):
            col, row = index % 2, index // 2
            kpis = seller.get("kpis") or {}
            opportunity = (seller.get("opportunities") or [{}])[0]
            message = opportunity.get("message") or "Sostener ejecución y seguimiento."
            self._card(
                slide, 0.65 + col * 6.05, 1.25 + row * 1.7, 5.7, 1.4,
                str(seller.get("identification", {}).get("name") or ""),
                f"Venta {self._money(kpis.get('sales'))} · crec. {self._number(kpis.get('growthPct'))}%\n{message}",
                accent=ACCENT if float(kpis.get("growthPct") or 0) >= 0 else WARN,
                value_size=10,
            )

    def _commitments(self, prs, report):
        slide = self._slide(prs, "Compromisos y próximos pasos", report)
        table = slide.shapes.add_table(
            6, 4, Inches(0.65), Inches(1.35), Inches(12.0), Inches(4.9)
        ).table
        headers = ["Acción", "Responsable", "Fecha", "Resultado esperado"]
        for col, label in enumerate(headers):
            self._table_cell(table.cell(0, col), label, GREEN, WHITE, True)
        for row in range(1, 6):
            for col in range(4):
                self._table_cell(table.cell(row, col), "", WHITE, DARK)
        self._textbox(slide, 0.7, 6.45, 11.8, 0.35, "Todo compromiso debe quedar asociado a evidencia y fecha de revisión.", 9, MUTED)

    def _slide(self, prs, title, report):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._background(slide, WHITE)
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.85))
        header.fill.solid()
        header.fill.fore_color.rgb = GREEN
        header.line.fill.background()
        self._textbox(slide, 0.55, 0.15, 10.8, 0.45, title, 23, WHITE, bold=True)
        self._textbox(
            slide, 10.75, 0.23, 2.0, 0.25,
            f"{report['period']['end']} · v{report['version']}", 8, WHITE, align=PP_ALIGN.RIGHT,
        )
        self._textbox(
            slide, 0.55, 7.12, 12.2, 0.18,
            f"Codenoa Sales Coach · snapshot {report['report_id'][:8]} · datos {report.get('data_updated_at') or 'sin fecha'}",
            7, MUTED,
        )
        return slide

    @staticmethod
    def _background(slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def _card(self, slide, x, y, width, height, label, value, accent=ACCENT, value_size=18):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y), Inches(width), Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT
        shape.line.color.rgb = accent
        self._textbox(slide, x + 0.18, y + 0.12, width - 0.35, 0.22, label, 9, accent, bold=True)
        self._textbox(slide, x + 0.18, y + 0.38, width - 0.35, height - 0.45, value, value_size, DARK, bold=value_size >= 16)

    def _bullet_box(self, slide, x, y, width, height, items, title):
        self._textbox(slide, x, y, width, 0.35, title, 14, GREEN, bold=True)
        text = "\n".join(f"• {str(item)}" for item in items[:6]) or "Sin información para este corte."
        self._textbox(slide, x, y + 0.4, width, height - 0.4, text, 12, DARK)

    def _empty(self, slide, message):
        self._textbox(slide, 1.0, 2.6, 11.3, 1.0, message, 20, MUTED, align=PP_ALIGN.CENTER)

    @staticmethod
    def _textbox(slide, x, y, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.text = str(text)
        paragraph.alignment = align
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        return box

    @staticmethod
    def _table_cell(cell, text, background, foreground, bold=False):
        cell.fill.solid()
        cell.fill.fore_color.rgb = background
        cell.text = str(text)
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(10)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = foreground
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)

    @staticmethod
    def _money(value):
        return f"$ {float(value or 0):,.0f}"

    @staticmethod
    def _number(value):
        return f"{float(value or 0):,.1f}"
