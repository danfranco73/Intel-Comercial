from __future__ import annotations

import mongomock
import pytest
from io import BytesIO
from pptx import Presentation

from sales_coach.repositories import MeetingRepository
from sales_coach.schemas.meetings import MeetingCreateRequest
from sales_coach.services.meeting_service import MeetingService
from sales_coach.services.pdf_service import MeetingPdfRenderer
from sales_coach.services.pptx_service import MeetingPptxRenderer
from security.models import AuthenticatedUser


class SalesCoachStub:
    def _report(self, start, end, scope):
        seller_rows = [
            {
                "sellerKey": "S1",
                "seller": "Ana",
                "salesForce": "Norte",
                "sales": 1200,
                "previousSales": 1000,
                "quantity": 60,
                "clients": 10,
                "avgTicket": 120,
                "growthPct": 20,
                "sharePct": 60,
                "mixCount": 4,
                "top3ClientsSharePct": 45,
                "rankSales": 1,
                "recoveredClients": 2,
                "droppedClients": 1,
                "strengths": [{"message": "Creció 20%.", "evidence": {"growthPct": 20}}],
                "opportunities": [{"message": "Ampliar mix.", "evidence": {"mix": 4}}],
                "actionPlan": [{"action": "Visitar cuentas prioritarias.", "evidence": {"clients": 10}}],
            },
            {
                "sellerKey": "S2",
                "seller": "Bruno",
                "salesForce": "Sur",
                "sales": 800,
                "previousSales": 1000,
                "quantity": 40,
                "clients": 8,
                "avgTicket": 100,
                "growthPct": -20,
                "sharePct": 40,
                "mixCount": 2,
                "top3ClientsSharePct": 60,
                "rankSales": 2,
                "recoveredClients": 0,
                "droppedClients": 2,
                "strengths": [],
                "opportunities": [],
                "actionPlan": [],
            },
        ]
        allowed = set(scope.get("seller_name") or [])
        if "seller_name" in scope:
            seller_rows = [row for row in seller_rows if row["seller"] in allowed]
        report = {
            "meta": {
                "comparison": {
                    "selectedStart": start,
                    "selectedEnd": end,
                    "comparisonStart": "2026-06-01",
                    "comparisonEnd": "2026-06-30",
                }
            },
            "dashboards": {
                "sellers": {"rows": seller_rows},
                "clients": {"summary": {"activeCount": sum(row["clients"] for row in seller_rows)}},
                "opportunities": {"summary": {"totalPotential": 350}, "rows": []},
                "history": {"charts": {"monthly": []}},
            },
        }
        return report, "fixture"

    @staticmethod
    def _freshness():
        return "2026-07-31T12:00:00"

    @staticmethod
    def _pct_change(current, previous):
        return round((current - previous) / previous * 100, 1) if previous else 0


def auth_user(role="admin", seller_key=None):
    return AuthenticatedUser(
        id=f"{role}-id",
        email=f"{role}@example.test",
        name=role,
        role=role,
        seller_key=seller_key,
    )


def test_request_rejects_unknown_filters_and_invalid_dates():
    with pytest.raises(ValueError, match="Filtro no permitido"):
        MeetingCreateRequest.parse(
            {
                "fechaDesde": "2026-07-01",
                "fechaHasta": "2026-07-31",
                "filters": {"unsafe": ["x"]},
            }
        )
    with pytest.raises(ValueError, match="posterior"):
        MeetingCreateRequest.parse(
            {"fechaDesde": "2026-08-01", "fechaHasta": "2026-07-31"}
        )


def test_snapshot_is_versioned_and_retrievable_without_recalculation():
    db = mongomock.MongoClient().db
    db.erp_sellers.insert_many(
        [
            {"seller_key": "S1", "seller_name": "Ana"},
            {"seller_key": "S2", "seller_name": "Bruno"},
        ]
    )
    service = MeetingService(db, SalesCoachStub())

    report = service.create(
        {
            "report_type": "general",
            "fechaDesde": "2026-07-01",
            "fechaHasta": "2026-07-31",
        },
        auth_user(),
        {},
    )
    persisted = service.get(report["report_id"], auth_user(), {})

    assert report["version"] == 1
    assert persisted["payload"]["kpis"]["net_sales"] == 2000
    assert persisted["data_updated_at"] == "2026-07-31T12:00:00"
    assert persisted["payload"]["executive_summary"] == report["payload"]["executive_summary"]


def test_seller_packet_is_restricted_to_own_scope():
    db = mongomock.MongoClient().db
    db.erp_sellers.insert_many(
        [
            {"seller_key": "S1", "seller_name": "Ana"},
            {"seller_key": "S2", "seller_name": "Bruno"},
        ]
    )
    service = MeetingService(db, SalesCoachStub())
    seller = auth_user("seller", "S1")

    report = service.create(
        {
            "report_type": "seller_packet",
            "fechaDesde": "2026-07-01",
            "fechaHasta": "2026-07-31",
        },
        seller,
        {"seller_name": ["Ana"]},
    )

    assert report["seller_keys"] == ["S1"]
    assert [page["identification"]["seller_key"] for page in report["payload"]["seller_pages"]] == ["S1"]


def test_supervisor_cannot_download_snapshot_with_foreign_seller():
    db = mongomock.MongoClient().db
    db.erp_sellers.insert_many(
        [
            {"seller_key": "S1", "seller_name": "Ana"},
            {"seller_key": "S2", "seller_name": "Bruno"},
        ]
    )
    repository = MeetingRepository(db)
    report = repository.create(
        {
            "company_key": "CODENOA",
            "report_type": "general",
            "period": {"start": "2026-07-01", "end": "2026-07-31"},
            "seller_keys": ["S2"],
            "payload": {},
        },
        "admin",
    )
    service = MeetingService(db, SalesCoachStub())

    with pytest.raises(PermissionError):
        service.get(
            report["report_id"],
            auth_user("supervisor"),
            {"seller_name": ["Ana"]},
        )


def test_pdf_is_valid_a4_document_and_contains_one_page_per_seller():
    report = {
        "report_id": "report-1",
        "version": 1,
        "report_type": "seller_packet",
        "period": {"start": "2026-07-01", "end": "2026-07-31"},
        "data_updated_at": "2026-07-31T12:00:00",
        "payload": {
            "title": "Codenoa Sales Coach",
            "kpis": {"net_sales": 2000, "active_clients": 18},
            "executive_summary": ["Venta verificada."],
            "ranking": [],
            "alerts": [],
            "recognitions": [],
            "seller_pages": [
                {
                    "identification": {"name": "Ana", "sales_force": "Norte", "period": "Julio"},
                    "kpis": {"sales": 1200, "clients": 10},
                    "strengths": [],
                    "opportunities": [],
                    "action_plan": [],
                    "objectives": [],
                    "alerts": [],
                },
                {
                    "identification": {"name": "Bruno", "sales_force": "Sur", "period": "Julio"},
                    "kpis": {"sales": 800, "clients": 8},
                    "strengths": [],
                    "opportunities": [],
                    "action_plan": [],
                    "objectives": [],
                    "alerts": [],
                },
            ],
        },
    }

    content = MeetingPdfRenderer().render(report)

    assert content.startswith(b"%PDF-")
    assert len(content) > 4000
    assert content.count(b"/Type /Page") >= 3


def test_pptx_is_editable_16_9_and_stays_below_twelve_slides():
    ranking = [
        {
            "seller": f"Vendedor {index}",
            "sales": 1000 - index * 50,
            "growthPct": 20 - index,
            "rankSales": index,
            "clients": 10 + index,
            "mixCount": 4,
        }
        for index in range(1, 11)
    ]
    report = {
        "report_id": "report-pptx",
        "version": 1,
        "report_type": "seller_packet",
        "period": {"start": "2026-07-01", "end": "2026-07-31"},
        "data_updated_at": "2026-07-31T12:00:00",
        "payload": {
            "title": "Codenoa Sales Coach",
            "kpis": {"net_sales": 2000, "active_clients": 18, "growth_pct": 5},
            "executive_summary": ["Venta verificada.", "Crecimiento positivo."],
            "ranking": ranking,
            "alerts": [
                {
                    "severity": "warning",
                    "type": "seller_decline",
                    "message": "Caída con evidencia.",
                }
            ],
            "recognitions": [{"seller": "Ana", "message": "Crecimiento 20%"}],
            "objectives": [],
            "seller_pages": [
                {
                    "identification": {"name": row["seller"]},
                    "kpis": row,
                    "opportunities": [],
                }
                for row in ranking
            ],
        },
    }

    content = MeetingPptxRenderer().render(report)
    presentation = Presentation(BytesIO(content))

    assert content.startswith(b"PK")
    assert 1 <= len(presentation.slides) <= 12
    assert presentation.slide_width / presentation.slide_height == pytest.approx(16 / 9, rel=0.001)
    assert any(shape.has_chart for slide in presentation.slides for shape in slide.shapes)
    assert any(shape.has_table for slide in presentation.slides for shape in slide.shapes)
